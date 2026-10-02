"""Integrity and regression checks on docs/proforma-reference.xlsx.

The workbook is an independent check on the pro-forma formulas, so these tests do not
recompute anything: they pin its layout, keep its assumptions equal to the Dallas pack, and
hold its recalculated values to figures computed by hand from the pro-forma formulas.
"""

from decimal import Decimal
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from feasibility.markets.loader import get_pack

WORKBOOK = Path(__file__).resolve().parents[1] / "docs" / "proforma-reference.xlsx"
SHEETS = ["Summary", "Assumptions", "S1", "S2", "S3", "Sensitivity"]
SCENARIOS = ["S1", "S2", "S3"]
KEY_COLUMN = 4  # D
VALUE_COLUMN = 2  # B
AT_MAX_VALUE_COLUMN = 5  # E
AT_MAX_KEY_COLUMN = 6  # F
# LibreOffice writes a boolean input back as a constant formula; nothing else may be a formula.
BOOLEAN_INPUT_FORMULAS = {"=TRUE()", "=FALSE()"}
CENT = Decimal("0.005")
RATIO = Decimal("0.00005")
EXACT_RATIO = RATIO / 10  # ratios in the workbook are rounded results, so compare tightly


@pytest.fixture(scope="module")
def formulas() -> Any:
    return load_workbook(WORKBOOK)


@pytest.fixture(scope="module")
def values() -> Any:
    return load_workbook(WORKBOOK, data_only=True)


def key_rows(sheet: Worksheet, key_column: int = KEY_COLUMN) -> dict[str, int]:
    rows: dict[str, int] = {}
    for row in range(2, sheet.max_row + 1):
        key = sheet.cell(row, key_column).value
        if key:
            assert key not in rows, f"{sheet.title}: duplicate key {key}"
            rows[key] = row
    return rows


def number(cell_value: object) -> Decimal:
    assert isinstance(cell_value, int | float | Decimal), f"not a number: {cell_value!r}"
    return Decimal(str(round(cell_value, 8)))


def test_the_workbook_has_the_sheets_in_the_layout_contract(formulas: Any) -> None:
    assert formulas.sheetnames == SHEETS


def test_the_workbook_has_no_hidden_sheets_or_defined_names(formulas: Any) -> None:
    assert all(sheet.sheet_state == "visible" for sheet in formulas)
    assert not list(formulas.defined_names)


def test_no_sheet_has_merged_cells_or_hidden_rows(formulas: Any) -> None:
    for sheet in formulas:
        assert not sheet.merged_cells.ranges, sheet.title
        assert not [row for row, dim in sheet.row_dimensions.items() if dim.hidden], sheet.title


# --- Assumptions equal the pack ------------------------------------------------------------

PACK_PATHS = {
    "as.acq_closing_pct": ("acquisition", "closing_pct"),
    "as.demo_flat": ("demolition", "flat"),
    "as.demo_per_sqft": ("demolition", "per_sqft"),
    "as.demo_fallback_sqft": ("demolition", "fallback_sqft"),
    "as.hard_cost_psf": ("construction", "hard_cost_per_sqft"),
    "as.contingency_pct": ("construction", "contingency_pct"),
    "as.soft_cost_pct": ("construction", "soft_cost_pct"),
    "as.build_share_pct": ("construction", "build_share_pct"),
    "as.ltc_pct": ("financing", "loan_to_cost_pct"),
    "as.rate_pct": ("financing", "rate_pct"),
    "as.points_pct": ("financing", "points_pct"),
    "as.draw_count": ("financing", "draw_count"),
    "as.draw_fee": ("financing", "draw_fee"),
    "as.hold_months": ("holding", "hold_months"),
    "as.tax_rate_pct": ("holding", "property_tax_rate_pct"),
    "as.insurance_pct": ("holding", "insurance_pct_of_hard_cost_per_year"),
    "as.commission_pct": ("selling", "commission_pct"),
    "as.sell_closing_pct": ("selling", "closing_pct"),
    "as.target_margin_pct": ("target", "margin_pct"),
    "as.min_home_sqft": ("sizing", "min_home_sqft"),
    "as.max_home_sqft": ("sizing", "max_home_sqft"),
    "as.min_comps": ("arv", "min_comps"),
    "as.min_comp_sqft": ("arv", "min_comp_sqft"),
    "as.new_build_premium_pct": ("arv", "new_build_premium_pct"),
}
RULE_FIELDS = ("coverage_pct", "stories", "living_share_pct")


def test_every_assumption_equals_the_dallas_pack(values: Any) -> None:
    sheet = values["Assumptions"]
    rows = key_rows(sheet)
    costs = get_pack("dallas").cost_assumptions
    expected: dict[str, Decimal] = {}
    for key, (group, field) in PACK_PATHS.items():
        expected[key] = Decimal(str(getattr(getattr(costs, group), field)))
    for zoning in ("R-7.5(A)", "default"):
        rule = costs.sizing.default if zoning == "default" else costs.sizing.rules[zoning]
        for field in RULE_FIELDS:
            expected[f"as.rule.{zoning}.{field}"] = Decimal(str(getattr(rule, field)))

    assert set(rows) == set(expected)
    for key, want in expected.items():
        assert number(sheet.cell(rows[key], VALUE_COLUMN).value) == want, key


# --- structure of the scenario sheets ------------------------------------------------------


def input_keys(keys: dict[str, int]) -> set[str]:
    return {key for key in keys if key.startswith("in.")}


@pytest.mark.parametrize("name", SCENARIOS)
def test_scenario_inputs_are_values_and_every_other_key_row_is_a_formula(
    formulas: Any, name: str
) -> None:
    sheet = formulas[name]
    rows = key_rows(sheet)
    inputs = input_keys(rows)
    assert {"in.price", "in.lot_sqft", "in.zoning", "in.is_vacant", "in.existing_sqft"} <= inputs
    assert "in.comp.1.price" in inputs
    assert "in.comp.1.area" in inputs

    for key, row in rows.items():
        cell = sheet.cell(row, VALUE_COLUMN).value
        is_formula = isinstance(cell, str) and cell.startswith("=")
        if key in inputs:
            assert not is_formula or cell in BOOLEAN_INPUT_FORMULAS, key
        else:
            assert is_formula, f"{name} {key} is not a formula: {cell!r}"


@pytest.mark.parametrize("name", SCENARIOS)
def test_scenario_formulas_reference_assumptions_not_retyped_constants(
    formulas: Any, name: str
) -> None:
    sheet = formulas[name]
    rows = key_rows(sheet)
    for key in ("cost.hard", "cost.contingency", "cost.soft", "fin.loan", "hold.tax"):
        assert "Assumptions!" in sheet.cell(rows[key], VALUE_COLUMN).value, key


@pytest.mark.parametrize("name", SCENARIOS)
def test_no_key_cell_is_an_error_and_only_documented_cells_are_blank(
    values: Any, name: str
) -> None:
    sheet = values[name]
    rows = key_rows(sheet)
    # Blank by design: S2/S3 have no existing area, S3 has no viable offer and so no headroom.
    blank_by_design = {"in.existing_sqft"}
    if name == "S3":
        blank_by_design |= {"max.headroom", "max.profit_at_max"}
    for key, row in rows.items():
        cell = sheet.cell(row, VALUE_COLUMN).value
        if cell is None:
            assert key in blank_by_design, f"{name} {key} is blank"
        if isinstance(cell, str):
            assert not cell.startswith("#"), f"{name} {key} is an error: {cell}"


def test_the_at_max_column_repeats_the_keys_with_an_atmax_prefix(formulas: Any) -> None:
    for name in ("S1", "S2"):
        sheet = formulas[name]
        at_max = key_rows(sheet, AT_MAX_KEY_COLUMN)
        assert "atmax.tot.profit" in at_max
        for key, row in at_max.items():
            assert key.startswith("atmax.")
            assert sheet.cell(row, AT_MAX_VALUE_COLUMN).value.startswith("=")


# --- hand-computed expected values, to the cent ------------------------------

MONEY: dict[str, tuple[str, str, str]] = {
    "sizing.uncapped": ("3168.00", "5016.33", "2480.00"),
    "sizing.buildable_sqft": ("3168", "3500", "2480"),
    "arv.arv": ("1420799.85", "1489400.50", "678417.76"),
    "cost.acq_closing": ("4200.00", "3000.00", "4950.00"),
    "cost.demolition": ("11700.00", "0.00", "14500.00"),
    "cost.hard": ("601920.00", "665000.00", "471200.00"),
    "cost.contingency": ("30096.00", "33250.00", "23560.00"),
    "cost.soft": ("72230.40", "79800.00", "56544.00"),
    "fin.financeable": ("1135946.40", "1078050.00", "1060804.00"),
    "fin.loan": ("908757.12", "862440.00", "848643.20"),
    "fin.interest_front": ("30235.82", "22788.00", "33962.64"),
    "fin.interest_progressive": ("24648.62", "27231.75", "19295.64"),
    "fin.points": ("18175.14", "17248.80", "16972.86"),
    "fin.total": ("74059.58", "68268.55", "71231.14"),
    "hold.tax": ("7014.14", "5010.10", "8266.66"),
    "hold.insurance": ("6771.60", "7481.25", "5301.00"),
    "sell.total": ("85247.99", "89364.04", "40705.07"),
    "tot.total_cost": ("1313239.71", "1251173.94", "1191257.87"),
    "tot.profit": ("107560.14", "238226.56", "-512840.11"),
    "tot.cash_invested": ("319234.60", "299369.90", "301909.60"),
}
RATIOS: dict[str, tuple[str, str, str]] = {
    "arv.median_psf": ("448.4848", "425.5430", "273.55555"),
    "tot.margin": ("0.0757", "0.1599", "-0.7559"),
    "tot.roi": ("0.3369", "0.7958", "-1.6987"),
    "tot.annualized": ("0.4492", "1.0610", "-2.2649"),
    "max.price_coeff": ("1.102700325", "1.102700325", "1.102700325"),
}
FIXED_PART = ("850105.5804", "920363.8400", "645421.2140")
MAX_OFFER = ("324271.50", "313436.54", None)  # S3 has no viable offer
HEADROOM = ("-95728.50", "13436.54", None)


def scenario_value(values: Any, name: str, key: str) -> Any:
    sheet = values[name]
    return sheet.cell(key_rows(sheet)[key], VALUE_COLUMN).value


@pytest.mark.parametrize("index", [0, 1, 2], ids=SCENARIOS)
def test_money_lines_match_the_formula_section_to_the_cent(values: Any, index: int) -> None:
    name = SCENARIOS[index]
    for key, expected in MONEY.items():
        got = number(scenario_value(values, name, key))
        assert abs(got - Decimal(expected[index])) <= CENT, f"{name} {key}: {got}"


@pytest.mark.parametrize("index", [0, 1, 2], ids=SCENARIOS)
def test_ratios_match_the_formula_section(values: Any, index: int) -> None:
    name = SCENARIOS[index]
    for key, expected in RATIOS.items():
        got = number(scenario_value(values, name, key))
        assert abs(got - Decimal(expected[index])) <= EXACT_RATIO, f"{name} {key}: {got}"


@pytest.mark.parametrize("index", [0, 1, 2], ids=SCENARIOS)
def test_the_maximum_offer_lines_match_the_formula_section(values: Any, index: int) -> None:
    name = SCENARIOS[index]
    fixed = number(scenario_value(values, name, "max.fixed_part"))
    assert abs(fixed - Decimal(FIXED_PART[index])) <= Decimal("0.0001")

    offer = scenario_value(values, name, "max.max_offer")
    headroom = scenario_value(values, name, "max.headroom")
    if MAX_OFFER[index] is None:
        assert offer == "none"
        assert headroom in (None, "")
        assert number(scenario_value(values, name, "max.numerator")) < 0
    else:
        assert abs(number(offer) - Decimal(str(MAX_OFFER[index]))) <= CENT
        assert abs(number(headroom) - Decimal(str(HEADROOM[index]))) <= CENT


@pytest.mark.parametrize(("name", "profit"), [("S1", "213119.98"), ("S2", "223410.09")])
def test_profit_at_the_maximum_offer_is_the_target_within_five_cents(
    values: Any, name: str, profit: str
) -> None:
    sheet = values[name]
    at_max = number(
        sheet.cell(
            key_rows(sheet, AT_MAX_KEY_COLUMN)["atmax.tot.profit"], AT_MAX_VALUE_COLUMN
        ).value
    )
    target = number(scenario_value(values, name, "max.target_profit"))

    assert abs(at_max - Decimal(profit)) <= CENT
    assert abs(at_max - target) <= Decimal("0.05")


def test_the_summary_shows_each_scenario_outcome(values: Any) -> None:
    rows = list(values["Summary"].iter_rows(min_row=2, max_row=4, values_only=True))

    assert [row[0] for row in rows] == SCENARIOS
    assert [number(row[1]) for row in rows] == [420000, 300000, 495000]
    for row, profit in zip(rows, ("107560.14", "238226.56", "-512840.11"), strict=True):
        assert abs(number(row[4]) - Decimal(profit)) <= CENT, row[0]
    assert rows[2][8] == "none"


# --- sensitivity grid ----------------------------------------------------------------------

GRID_SPOT_CHECKS = {
    (0, 0, 9): ("107560.14", "0.0757"),
    (10, -10, 6): ("337940.43", "0.2162"),
    (-10, 20, 12): ("-201522.39", "-0.1576"),
    (-5, 0, 6): ("63672.60", "0.0472"),
    (5, 20, 12): ("-1189.62", "-0.0008"),
}


def test_the_sensitivity_grid_has_sixty_rows_in_the_documented_order(values: Any) -> None:
    rows = list(values["Sensitivity"].iter_rows(min_row=2, max_row=61, max_col=9, values_only=True))
    grid = get_pack("dallas").cost_assumptions.sensitivity
    expected = [
        (d, e, h)
        for d in grid.arv_delta_pct
        for e in grid.hard_cost_delta_pct
        for h in grid.hold_months
    ]

    assert [tuple(row[:3]) for row in rows] == expected
    header = [cell.value for cell in values["Sensitivity"][1][:9]]
    assert header == [
        "arv_delta_pct",
        "hard_cost_delta_pct",
        "hold_months",
        "arv",
        "hard_cost",
        "total_cost",
        "profit",
        "margin",
        "roi",
    ]
    assert values["Sensitivity"].max_row == 61
    assert all(cell is not None for row in rows for cell in row)


def test_the_grid_spot_checks_and_its_centre_equal_the_base_case(values: Any) -> None:
    grid = {
        tuple(row[:3]): row
        for row in values["Sensitivity"].iter_rows(
            min_row=2, max_row=61, max_col=9, values_only=True
        )
    }
    for point, (profit, margin) in GRID_SPOT_CHECKS.items():
        assert abs(number(grid[point][6]) - Decimal(profit)) <= CENT, point
        assert abs(number(grid[point][7]) - Decimal(margin)) <= RATIO, point

    base_profit = number(scenario_value(values, "S1", "tot.profit"))
    base_margin = number(scenario_value(values, "S1", "tot.margin"))
    assert number(grid[(0, 0, 9)][6]) == base_profit
    assert number(grid[(0, 0, 9)][7]) == base_margin


def test_the_sensitivity_formulas_are_live(formulas: Any) -> None:
    rows = formulas["Sensitivity"].iter_rows(min_row=2, max_row=61, min_col=4, max_col=9)
    for row in rows:
        assert all(isinstance(cell.value, str) and cell.value.startswith("=") for cell in row)
