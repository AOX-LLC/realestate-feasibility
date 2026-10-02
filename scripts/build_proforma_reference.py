import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

MONEY = "#,##0.00"
RATIO = "0.0000"
ASSUME = [
    ("as.acq_closing_pct", 1.0, "Acquisition closing costs, percent of price"),
    ("as.demo_flat", 2500, "Demolition flat amount"),
    ("as.demo_per_sqft", 8, "Demolition per existing sq ft"),
    ("as.demo_fallback_sqft", 1500, "Existing sq ft assumed when unknown"),
    ("as.hard_cost_psf", 190, "Hard cost per buildable sq ft"),
    ("as.contingency_pct", 5, "Contingency, percent of hard cost"),
    ("as.soft_cost_pct", 12, "Soft costs, percent of hard cost"),
    ("as.build_share_pct", 70, "Share of hold spent building, percent"),
    ("as.ltc_pct", 80, "Loan to cost, percent"),
    ("as.rate_pct", 10.0, "Annual interest rate, percent"),
    ("as.points_pct", 2.0, "Loan points, percent of loan"),
    ("as.draw_count", 5, "Number of construction draws"),
    ("as.draw_fee", 200, "Fee per draw"),
    ("as.hold_months", 9, "Total hold in months"),
    ("as.tax_rate_pct", 2.226710, "Annual property tax rate, percent of price"),
    ("as.insurance_pct", 1.5, "Annual insurance, percent of hard cost"),
    ("as.commission_pct", 5.0, "Sell-side commission, percent of ARV"),
    ("as.sell_closing_pct", 1.0, "Sell-side closing, percent of ARV"),
    ("as.target_margin_pct", 15, "Target margin, percent of ARV"),
    ("as.min_home_sqft", 1500, "Minimum buildable home size"),
    ("as.max_home_sqft", 3500, "Maximum buildable home size"),
    ("as.min_comps", 3, "Minimum comps for an ARV"),
    ("as.min_comp_sqft", 600, "Comps below this area are ignored"),
    ("as.new_build_premium_pct", 0, "New-build premium on median price per sq ft, percent"),
    ("as.rule.R-7.5(A).coverage_pct", 45, "Zoning R-7.5(A): lot coverage percent"),
    ("as.rule.R-7.5(A).stories", 2, "Zoning R-7.5(A): stories"),
    ("as.rule.R-7.5(A).living_share_pct", 55, "Zoning R-7.5(A): living share of gross, percent"),
    ("as.rule.default.coverage_pct", 40, "Default zoning: lot coverage percent"),
    ("as.rule.default.stories", 2, "Default zoning: stories"),
    ("as.rule.default.living_share_pct", 50, "Default zoning: living share of gross, percent"),
]
AROW = {k: i + 2 for i, (k, _, _) in enumerate(ASSUME)}
ALAST = len(ASSUME) + 1


def assumption_ref(k):
    return f"Assumptions!$B${AROW[k]}"


SCEN = {
    "S1": dict(
        price=420000,
        lot=6400,
        zoning="R-7.5(A)",
        vacant=False,
        existing=1150,
        comps=[
            (1330000, 3000),
            (1420000, 3150),
            (1250000, 2800),
            (1480000, 3300),
            (1190000, 2650),
            (1390000, 3100),
            (1520000, 3350),
        ],
    ),
    "S2": dict(
        price=300000,
        lot=10134,
        zoning="R-7.5(A)",
        vacant=True,
        existing=None,
        comps=[(1450000, 3400), (1380000, 3250), (1500000, 3600), (1280000, 3000)],
    ),
    "S3": dict(
        price=495000,
        lot=6200,
        zoning="CD-12",
        vacant=False,
        existing=None,
        comps=[(640000, 2300), (690000, 2500), (610000, 2250), (720000, 2700)],
    ),
}


def rule(field, zk):
    f = f'"as.rule."&{{{zk}}}&".{field}"'
    rng_d = f"Assumptions!$D$2:$D${ALAST}"
    rng_b = f"Assumptions!$B$2:$B${ALAST}"
    return (
        f"=IF(ISNUMBER(MATCH({f},{rng_d},0)),INDEX({rng_b},MATCH({f},{rng_d},0)),"
        f'INDEX({rng_b},MATCH("as.rule.default.{field}",{rng_d},0)))'
    )


L_ = "{as.ltc_pct}/100"
I_ = "{as.rate_pct}/100/12"
H_ = "{as.hold_months}"


# (key, label, note, formula, fmt, in_at_max)
def lines(ncomps):
    c1, cn = "{COMPFIRST}", "{COMPLAST}"
    return [
        (
            "sizing.rule_coverage_pct",
            "Rule: coverage %",
            "Zoning rule row, or default when none",
            rule("coverage_pct", "in.zoning"),
            "0.00",
            0,
        ),
        (
            "sizing.rule_stories",
            "Rule: stories",
            "Zoning rule row, or default when none",
            rule("stories", "in.zoning"),
            "0",
            0,
        ),
        (
            "sizing.rule_living_share_pct",
            "Rule: living share %",
            "Zoning rule row, or default when none",
            rule("living_share_pct", "in.zoning"),
            "0.00",
            0,
        ),
        (
            "sizing.footprint",
            "Footprint sq ft",
            "Lot x coverage % / 100",
            "={in.lot_sqft}*{sizing.rule_coverage_pct}/100",
            "#,##0.00",
            0,
        ),
        (
            "sizing.gross",
            "Gross sq ft",
            "Footprint x stories",
            "={sizing.footprint}*{sizing.rule_stories}",
            "#,##0.00",
            0,
        ),
        (
            "sizing.uncapped",
            "Uncapped living sq ft",
            "Gross x living share % / 100",
            "={sizing.gross}*{sizing.rule_living_share_pct}/100",
            "#,##0.00",
            0,
        ),
        (
            "sizing.buildable_sqft",
            "Buildable sq ft (Q)",
            "ROUND(MIN(MAX(uncapped, min), max), 0)",
            "=ROUND(MIN(MAX({sizing.uncapped},{as.min_home_sqft}),{as.max_home_sqft}),0)",
            "#,##0",
            0,
        ),
        (
            "arv.comp_count_used",
            "Comps used",
            "Count of comps with area >= min comp sq ft",
            f"=COUNT(G{c1}:G{cn})",
            "0",
            0,
        ),
        (
            "arv.median_psf",
            "Median price per sq ft",
            "Median of psf_i helper column (column G)",
            f"=MEDIAN(G{c1}:G{cn})",
            "0.00000",
            0,
        ),
        (
            "arv.arv",
            "ARV (V)",
            "ROUND(median psf x Q x (1 + premium %/100), 2)",
            "=ROUND({arv.median_psf}*{sizing.buildable_sqft}*(1+{as.new_build_premium_pct}/100),2)",
            MONEY,
            0,
        ),
        (
            "cost.acq_closing",
            "Acquisition closing",
            "ROUND(P x closing % / 100, 2)",
            "=ROUND({in.price}*{as.acq_closing_pct}/100,2)",
            MONEY,
            1,
        ),
        (
            "cost.demolition",
            "Demolition (D)",
            "0 if vacant, else ROUND(flat + per sq ft x existing or fallback, 2)",
            "=IF({in.is_vacant},0,ROUND({as.demo_flat}+{as.demo_per_sqft}*IF(ISBLANK({in.existing_sqft}),{as.demo_fallback_sqft},{in.existing_sqft}),2))",
            MONEY,
            1,
        ),
        (
            "cost.hard",
            "Hard cost (K)",
            "ROUND(Q x hard cost per sq ft, 2)",
            "=ROUND({sizing.buildable_sqft}*{as.hard_cost_psf},2)",
            MONEY,
            1,
        ),
        (
            "cost.contingency",
            "Contingency (Kc)",
            "ROUND(K x contingency % / 100, 2)",
            "=ROUND({cost.hard}*{as.contingency_pct}/100,2)",
            MONEY,
            1,
        ),
        (
            "cost.soft",
            "Soft costs (S)",
            "ROUND(K x soft % / 100, 2)",
            "=ROUND({cost.hard}*{as.soft_cost_pct}/100,2)",
            MONEY,
            1,
        ),
        (
            "fin.months_construction",
            "Months of construction (C)",
            "H x build share % / 100",
            "={as.hold_months}*{as.build_share_pct}/100",
            "0.00",
            1,
        ),
        (
            "fin.months_sale",
            "Months on sale (Mo)",
            "H - C",
            "={as.hold_months}-{fin.months_construction}",
            "0.00",
            1,
        ),
        (
            "fin.financeable",
            "Financeable cost (F)",
            "P + D + K + Kc + S",
            "={in.price}+{cost.demolition}+{cost.hard}+{cost.contingency}+{cost.soft}",
            MONEY,
            1,
        ),
        (
            "fin.loan",
            "Loan (Ln)",
            "ROUND(F x LTC % / 100, 2)",
            "=ROUND({fin.financeable}*{as.ltc_pct}/100,2)",
            MONEY,
            1,
        ),
        (
            "fin.interest_front",
            "Interest, front-drawn",
            "ROUND((P + D + S) x l x i x H, 2)",
            f"=ROUND(({{in.price}}+{{cost.demolition}}+{{cost.soft}})*{L_}*{I_}*{H_},2)",
            MONEY,
            1,
        ),
        (
            "fin.interest_progressive",
            "Interest, progressive",
            "ROUND((K + Kc) x l x i x (C/2 + Mo), 2)",
            f"=ROUND(({{cost.hard}}+{{cost.contingency}})*{L_}*{I_}*({{fin.months_construction}}/2+{{fin.months_sale}}),2)",
            MONEY,
            1,
        ),
        (
            "fin.interest",
            "Interest",
            "Front + progressive",
            "={fin.interest_front}+{fin.interest_progressive}",
            MONEY,
            1,
        ),
        (
            "fin.points",
            "Points",
            "ROUND(Ln x points % / 100, 2)",
            "=ROUND({fin.loan}*{as.points_pct}/100,2)",
            MONEY,
            1,
        ),
        (
            "fin.draw_fees",
            "Draw fees",
            "Draw count x draw fee",
            "={as.draw_count}*{as.draw_fee}",
            MONEY,
            1,
        ),
        (
            "fin.total",
            "Financing total",
            "Interest + points + draw fees",
            "={fin.interest}+{fin.points}+{fin.draw_fees}",
            MONEY,
            1,
        ),
        (
            "hold.tax",
            "Holding: property tax",
            "ROUND(P x tax % / 100 x H / 12, 2)",
            "=ROUND({in.price}*{as.tax_rate_pct}/100*{as.hold_months}/12,2)",
            MONEY,
            1,
        ),
        (
            "hold.insurance",
            "Holding: insurance",
            "ROUND(K x insurance % / 100 x H / 12, 2)",
            "=ROUND({cost.hard}*{as.insurance_pct}/100*{as.hold_months}/12,2)",
            MONEY,
            1,
        ),
        (
            "hold.total",
            "Holding total",
            "Tax + insurance",
            "={hold.tax}+{hold.insurance}",
            MONEY,
            1,
        ),
        (
            "sell.commission",
            "Selling: commission",
            "ROUND(V x commission % / 100, 2)",
            "=ROUND({arv.arv}*{as.commission_pct}/100,2)",
            MONEY,
            1,
        ),
        (
            "sell.closing",
            "Selling: closing",
            "ROUND(V x sell closing % / 100, 2)",
            "=ROUND({arv.arv}*{as.sell_closing_pct}/100,2)",
            MONEY,
            1,
        ),
        (
            "sell.total",
            "Selling total",
            "Commission + closing",
            "={sell.commission}+{sell.closing}",
            MONEY,
            1,
        ),
        (
            "tot.total_cost",
            "Total cost",
            "P + closing + D + K + Kc + S + financing + holding + selling",
            "={in.price}+{cost.acq_closing}+{cost.demolition}+{cost.hard}+{cost.contingency}+{cost.soft}+{fin.total}+{hold.total}+{sell.total}",
            MONEY,
            1,
        ),
        ("tot.profit", "Profit", "V - total cost", "={arv.arv}-{tot.total_cost}", MONEY, 1),
        (
            "tot.margin",
            "Margin",
            "ROUND(profit / V, 4)",
            "=ROUND({tot.profit}/{arv.arv},4)",
            RATIO,
            1,
        ),
        (
            "tot.cash_invested",
            "Cash invested",
            "Total cost - selling total - loan",
            "={tot.total_cost}-{sell.total}-{fin.loan}",
            MONEY,
            1,
        ),
        (
            "tot.roi",
            "ROI",
            "ROUND(profit / cash invested, 4); blank if cash <= 0",
            '=IF({tot.cash_invested}<=0,"",ROUND({tot.profit}/{tot.cash_invested},4))',
            RATIO,
            1,
        ),
        (
            "tot.annualized",
            "Annualized ROI",
            "ROUND(profit / cash x 12 / H, 4), simple",
            '=IF({tot.cash_invested}<=0,"",ROUND({tot.profit}/{tot.cash_invested}*12/{as.hold_months},4))',
            RATIO,
            1,
        ),
        (
            "max.fixed_part",
            "Max offer: fixed part (A)",
            "D+K+Kc+S + l*pts*(D+K+Kc+S) + l*i*((D+S)*H+(K+Kc)*(C/2+Mo))"
            " + draw fees + insurance + selling",
            f"={{cost.demolition}}+{{cost.hard}}+{{cost.contingency}}+{{cost.soft}}+{L_}*({{as.points_pct}}/100)*({{cost.demolition}}+{{cost.hard}}+{{cost.contingency}}+{{cost.soft}})+{L_}*{I_}*(({{cost.demolition}}+{{cost.soft}})*{H_}+({{cost.hard}}+{{cost.contingency}})*({{fin.months_construction}}/2+{{fin.months_sale}}))+{{fin.draw_fees}}+{{hold.insurance}}+{{sell.total}}",
            "#,##0.0000",
            0,
        ),
        (
            "max.price_coeff",
            "Max offer: price coefficient (B)",
            "1 + closing % + l*(points % + i*H) + tax % x H/12",
            f"=1+{{as.acq_closing_pct}}/100+{L_}*({{as.points_pct}}/100+{I_}*{H_})+({{as.tax_rate_pct}}/100)*{H_}/12",
            "0.0000000000",
            0,
        ),
        (
            "max.numerator",
            "Max offer: numerator (N)",
            "V x (1 - target margin) - A",
            "={arv.arv}*(1-{as.target_margin_pct}/100)-{max.fixed_part}",
            "#,##0.0000",
            0,
        ),
        (
            "max.max_offer",
            "Maximum offer",
            "none if N < 0, else ROUNDDOWN(N / B, 2)",
            '=IF({max.numerator}<0,"none",ROUNDDOWN({max.numerator}/{max.price_coeff},2))',
            MONEY,
            0,
        ),
        (
            "max.headroom",
            "Headroom vs list price",
            "Max offer - P; blank when none",
            '=IF(ISNUMBER({max.max_offer}),{max.max_offer}-{in.price},"")',
            MONEY,
            0,
        ),
        (
            "max.target_profit",
            "Target profit",
            "ROUND(V x target margin, 2)",
            "=ROUND({arv.arv}*{as.target_margin_pct}/100,2)",
            MONEY,
            0,
        ),
        (
            "max.profit_at_max",
            "Profit at max offer",
            "Profit from the at-max-offer column (E)",
            '=IF(ISNUMBER({max.max_offer}),{atmax.tot.profit},"")',
            MONEY,
            0,
        ),
    ]


hdr_font = Font(bold=True)
hdr_fill = PatternFill("solid", fgColor="DDDDDD")


def header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths, strict=True), 1):
        cell = ws.cell(1, i, c)
        cell.font = hdr_font
        cell.fill = hdr_fill
        ws.column_dimensions[cell.column_letter].width = w
    ws.freeze_panes = "A2"


wb = Workbook()
ws_summary = wb.active
ws_summary.title = "Summary"
ws_assumptions = wb.create_sheet("Assumptions")
header(ws_assumptions, ["Label", "Value", "Note", "Key"], [38, 14, 60, 36])
for k, (key, val, note) in enumerate(ASSUME):
    r = k + 2
    ws_assumptions.cell(r, 1, key.split("as.", 1)[1].replace("_", " ").replace(".", " "))
    ws_assumptions.cell(r, 2, val)
    ws_assumptions.cell(r, 3, note)
    ws_assumptions.cell(r, 4, key)

ROWS = {}
for name, sc in SCEN.items():
    ws = wb.create_sheet(name)
    header(
        ws,
        ["Label", "Value", "Note", "Key", "At max offer", "At-max key", "psf_i helper"],
        [34, 18, 62, 34, 18, 34, 14],
    )
    r = 2
    ref = {}

    def put(label, val, note, key, fmt=None, ws=ws, ref=ref):
        global r
        ws.cell(r, 1, label)
        ws.cell(r, 2, val)
        ws.cell(r, 3, note)
        ws.cell(r, 4, key)
        if fmt:
            ws.cell(r, 2).number_format = fmt
        ref[key] = r
        r += 1

    put("List price (P)", sc["price"], "Input", "in.price", MONEY)
    put("Lot sq ft (L)", sc["lot"], "Input", "in.lot_sqft", "#,##0")
    put("Zoning", sc["zoning"], "Input; no rule row means default", "in.zoning")
    put("Vacant", sc["vacant"], "Input", "in.is_vacant")
    put("Existing sq ft", sc["existing"], "Input; blank means unknown", "in.existing_sqft", "#,##0")
    first = r
    for n, (p, a) in enumerate(sc["comps"], 1):
        put(f"Comp {n} price", p, "Input", f"in.comp.{n}.price", MONEY)
        put(
            f"Comp {n} area",
            a,
            "Input; used only if area >= min comp sq ft",
            f"in.comp.{n}.area",
            "#,##0",
        )
        ws.cell(
            r - 1,
            7,
            f'=IF(B{r - 1}>={assumption_ref("as.min_comp_sqft")},ROUND(B{r - 2}/B{r - 1},4),"")',
        ).number_format = RATIO
    last = r - 1
    L = lines(len(sc["comps"]))
    # assign rows first
    start = r
    for i, entry in enumerate(L):
        ref[entry[0]] = start + i
    emap = {entry[0] for entry in L if entry[5]}
    for entry in L:
        ref["atmax." + entry[0]] = ref[entry[0]]
    maxrow = ref["max.max_offer"]

    def render(tpl, col, first=first, last=last, ref=ref, maxrow=maxrow, emap=emap):
        def sub(m):
            k = m.group(1)
            if k == "COMPFIRST":
                return str(first)
            if k == "COMPLAST":
                return str(last)
            if k.startswith("as."):
                return assumption_ref(k)
            if k.startswith("atmax."):
                return f"E{ref[k]}"
            if col == "E" and k == "in.price":
                return f"$B${maxrow}"
            if col == "E" and k in emap:
                return f"E{ref[k]}"
            return f"B{ref[k]}"

        return re.sub(r"\{([^}]+)\}", sub, tpl)

    for key, label, note, f, fmt, in_at_max in L:
        ws.cell(r, 1, label)
        ws.cell(r, 3, note)
        ws.cell(r, 4, key)
        ws.cell(r, 2, render(f, "B")).number_format = fmt
        if in_at_max:
            body = render(f, "E")[1:]
            ws.cell(r, 5, f'=IF(ISNUMBER($B${maxrow}),{body},"")').number_format = fmt
            ws.cell(r, 6, "atmax." + key)
        r += 1
    assert r - 1 == start + len(L) - 1  # noqa: S101
    ROWS[name] = ref
    ws.column_dimensions["G"].width = 14

# Sensitivity
ws_sensitivity = wb.create_sheet("Sensitivity")
cols = [
    "arv_delta_pct",
    "hard_cost_delta_pct",
    "hold_months",
    "arv",
    "hard_cost",
    "total_cost",
    "profit",
    "margin",
    "roi",
    "sell_commission",
    "sell_closing",
    "sell_total",
    "contingency",
    "soft",
    "insurance",
    "months_construction",
    "months_sale",
    "financeable",
    "loan",
    "interest_front",
    "interest_progressive",
    "interest",
    "points",
    "draw_fees",
    "fin_total",
    "tax",
    "hold_total",
    "cash_invested",
]
header(ws_sensitivity, cols, [14, 18, 12, 16, 14, 16, 16, 10, 10] + [16] * (len(cols) - 9))
R1 = ROWS["S1"]


def s1(k):
    return f"S1!$B${R1[k]}"


r = 2
for d in (-10, -5, 0, 5, 10):
    for e in (-10, 0, 10, 20):
        for h in (6, 9, 12):
            P = s1("in.price")
            f = {
                "D": f"=ROUND({s1('arv.arv')}*(1+A{r}/100),2)",
                "E": f"=ROUND({s1('cost.hard')}*(1+B{r}/100),2)",
                "J": f"=ROUND(D{r}*{assumption_ref('as.commission_pct')}/100,2)",
                "K": f"=ROUND(D{r}*{assumption_ref('as.sell_closing_pct')}/100,2)",
                "L": f"=J{r}+K{r}",
                "M": f"=ROUND(E{r}*{assumption_ref('as.contingency_pct')}/100,2)",
                "N": f"=ROUND(E{r}*{assumption_ref('as.soft_cost_pct')}/100,2)",
                "O": f"=ROUND(E{r}*{assumption_ref('as.insurance_pct')}/100*C{r}/12,2)",
                "P": f"=C{r}*{assumption_ref('as.build_share_pct')}/100",
                "Q": f"=C{r}-P{r}",
                "R": f"={P}+{s1('cost.demolition')}+E{r}+M{r}+N{r}",
                "S": f"=ROUND(R{r}*{assumption_ref('as.ltc_pct')}/100,2)",
                "T": (
                    f"=ROUND(({P}+{s1('cost.demolition')}+N{r})"
                    f"*{assumption_ref('as.ltc_pct')}/100"
                    f"*{assumption_ref('as.rate_pct')}/100/12*C{r},2)"
                ),
                "U": (
                    f"=ROUND((E{r}+M{r})*{assumption_ref('as.ltc_pct')}/100"
                    f"*{assumption_ref('as.rate_pct')}/100/12*(P{r}/2+Q{r}),2)"
                ),
                "V": f"=T{r}+U{r}",
                "W": f"=ROUND(S{r}*{assumption_ref('as.points_pct')}/100,2)",
                "X": f"={assumption_ref('as.draw_count')}*{assumption_ref('as.draw_fee')}",
                "Y": f"=V{r}+W{r}+X{r}",
                "Z": f"=ROUND({P}*{assumption_ref('as.tax_rate_pct')}/100*C{r}/12,2)",
                "AA": f"=Z{r}+O{r}",
                "F": (
                    f"={P}+{s1('cost.acq_closing')}+{s1('cost.demolition')}"
                    f"+E{r}+M{r}+N{r}+Y{r}+AA{r}+L{r}"
                ),
                "G": f"=D{r}-F{r}",
                "H": f"=ROUND(G{r}/D{r},4)",
                "AB": f"=F{r}-L{r}-S{r}",
                "I": f'=IF(AB{r}<=0,"",ROUND(G{r}/AB{r},4))',
            }
            ws_sensitivity.cell(r, 1, d)
            ws_sensitivity.cell(r, 2, e)
            ws_sensitivity.cell(r, 3, h)
            for c, v in f.items():
                cell = ws_sensitivity[f"{c}{r}"]
                cell.value = v
                cell.number_format = RATIO if c in ("H", "I") else MONEY
            r += 1

# Summary
sc_cols = [
    "Scenario",
    "Price",
    "ARV",
    "Total cost",
    "Profit",
    "Margin",
    "ROI",
    "Annualized ROI",
    "Max offer",
    "Headroom",
]
header(ws_summary, sc_cols, [10, 16, 16, 16, 16, 10, 10, 14, 16, 16])
for i, name in enumerate(SCEN, 2):
    R = ROWS[name]
    items = [
        ("in.price", MONEY),
        ("arv.arv", MONEY),
        ("tot.total_cost", MONEY),
        ("tot.profit", MONEY),
        ("tot.margin", RATIO),
        ("tot.roi", RATIO),
        ("tot.annualized", RATIO),
        ("max.max_offer", MONEY),
        ("max.headroom", MONEY),
    ]
    ws_summary.cell(i, 1, name)
    for j, (k, fmt) in enumerate(items, 2):
        ws_summary.cell(i, j, f"={name}!B{R[k]}").number_format = fmt
wb.save("proforma-reference.xlsx")
with Path("rows.json").open("w") as rows_file:
    json.dump(ROWS, rows_file)
